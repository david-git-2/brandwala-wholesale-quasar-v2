#!/usr/bin/env python3
"""Export WTS Excel catalog to wts_data.json."""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime, timezone

import openpyxl

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
ROOT_ENV_FILE = os.path.join(ROOT_DIR, ".env")
WEB_ENV_FILE = os.path.join(ROOT_DIR, "web", ".env")
sys.path.insert(0, os.path.join(ROOT_DIR, "python"))
from wts_excel_spec import REQUIRED_WTS_COLUMNS, all_field_candidates, sanitize_cell_text


def load_env_file(path: str) -> None:
    if not os.path.exists(path):
        return
    with open(path, "r", encoding="utf-8") as handle:
        for raw_line in handle:
            line = raw_line.strip()
            if not line or line.startswith("#"):
                continue
            if line.startswith("export "):
                line = line[7:].strip()
            if "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip()
            if (
                len(value) >= 2
                and ((value.startswith('"') and value.endswith('"')) or (value.startswith("'") and value.endswith("'")))
            ):
                value = value[1:-1]
            if key:
                os.environ.setdefault(key, value)


load_env_file(WEB_ENV_FILE)
load_env_file(ROOT_ENV_FILE)


def env_path(env_key: str, default_abs_path: str) -> str:
    raw = str(os.getenv(env_key, "")).strip()
    if not raw:
        return default_abs_path
    return raw if os.path.isabs(raw) else os.path.join(ROOT_DIR, raw)


DEFAULT_XLSX = env_path("PY_UK_WTS_XLSX_PATH", os.path.join(ROOT_DIR, "python", "data", "uk", "wts_data.xlsx"))
DEFAULT_OUT_JSON = env_path("PY_UK_WTS_OUT_JSON_PATH", os.path.join(ROOT_DIR, "web", "public", "uk", "wts_data.json"))
DEFAULT_OUT_MANIFEST = env_path(
    "PY_UK_WTS_OUT_MANIFEST_PATH",
    os.path.join(ROOT_DIR, "web", "public", "uk", "wts_manifest.json"),
)
DEFAULT_HEADER_ROW = 1


def log(msg: str) -> None:
    print(msg, flush=True)


def normalize_header(h: str) -> str:
    text = str(h).strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text


def resolve_header_name(header_to_col: dict, headers: list, candidates: list) -> str:
    for candidate in candidates:
        col = header_to_col.get(normalize_header(candidate))
        if col:
            return headers[col - 1]
    return ""


def to_text(value) -> str:
    return str(value if value is not None else "").strip()


def to_float_or_default(value, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(",", "")
    text = re.sub(r"[^0-9.\-]+", "", text)
    if text in ("", "-", ".", "-."):
        return default
    try:
        return float(text)
    except ValueError:
        return default


def to_int_or_default(value, default: int = 0) -> int:
    return int(round(to_float_or_default(value, float(default))))


def prompt_int(label: str, default: int) -> int:
    if not sys.stdin.isatty():
        return default
    while True:
        raw = input(f"{label} [{default}]: ").strip()
        if raw == "":
            return default
        try:
            val = int(raw)
            if val <= 0:
                log("Please enter a positive integer.")
                continue
            return val
        except ValueError:
            log("Invalid number. Try again.")


def format_duration(seconds: float) -> str:
    seconds = max(0, int(seconds))
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    if h > 0:
        return f"{h}h {m}m {s}s"
    if m > 0:
        return f"{m}m {s}s"
    return f"{s}s"


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Export WTS Excel catalog to JSON.")
    parser.add_argument("--header-row", type=int, default=None, help="1-based header row (skips prompt)")
    return parser.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    t0 = time.perf_counter()

    excel_path = DEFAULT_XLSX
    out_json_path = DEFAULT_OUT_JSON

    required_header_lines = "\n".join(f"    - {col['excel']} → {col['db']}" for col in REQUIRED_WTS_COLUMNS)
    log(
        "\n"
        "WTS export starting...\n"
        "\n"
        "Required headers (header row). Empty columns between them are OK:\n"
        f"{required_header_lines}\n"
        "\n"
        "Product ID: product_id = barcode + '_' + product_code\n"
    )

    if args.header_row is not None:
        if args.header_row <= 0:
            raise SystemExit("--header-row must be a positive integer")
        header_row = args.header_row
    else:
        header_row = prompt_int("Enter header row number", DEFAULT_HEADER_ROW)

    log(f"\nExcel: {excel_path}")
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel not found: {excel_path}")

    os.makedirs(os.path.dirname(out_json_path), exist_ok=True)
    log(f"Output JSON: {out_json_path}\n")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sh = wb[wb.sheetnames[0]]
    log(f"Using sheet: {sh.title}")

    max_col = sh.max_column
    max_row = sh.max_row
    log(f"Sheet size: rows={max_row}, cols={max_col}")

    log(f"Reading headers from row {header_row}...")
    headers = []
    for c in range(1, max_col + 1):
        v = sh.cell(header_row, c).value
        headers.append(str(v).strip() if v is not None else f"col_{c}")

    header_to_col = {}
    for idx, h in enumerate(headers, start=1):
        header_to_col.setdefault(normalize_header(h), idx)

    resolved_headers = {
        field: resolve_header_name(header_to_col, headers, candidates)
        for field, candidates in all_field_candidates().items()
    }

    missing_required = [col["excel"] for col in REQUIRED_WTS_COLUMNS if not resolved_headers.get(col["key"])]
    if missing_required:
        raise RuntimeError(
            f"Missing required header(s) in row {header_row}: {', '.join(missing_required)}\n"
            "Required headers: " + ", ".join(col["excel"] for col in REQUIRED_WTS_COLUMNS)
        )

    product_code_header = resolved_headers["product_code"]
    barcode_header = resolved_headers["barcode"]
    name_header = resolved_headers["name"]
    price_header = resolved_headers["price"]
    case_size_header = resolved_headers["case_size"]
    available_header = resolved_headers["available_units"]

    log(f"Found product_code column: {product_code_header}")
    log(f"Found barcode column: {barcode_header}")
    log(f"Found name column: {name_header}")
    log(f"Found price column: {price_header}")
    log(f"Found pack column: {case_size_header}")
    log(f"Found available column: {available_header}\n")

    start_data_row = header_row + 1
    products_by_row: dict[int, dict] = {}
    for r in range(start_data_row, max_row + 1):
        row_vals = [sh.cell(r, c).value for c in range(1, max_col + 1)]
        if all(v is None or str(v).strip() == "" for v in row_vals):
            continue
        obj = {headers[c - 1]: row_vals[c - 1] for c in range(1, max_col + 1)}
        obj["_rowNumber"] = r
        products_by_row[r] = obj

    log(f"Rows loaded: {len(products_by_row)}")

    required_value_keys = [col["key"] for col in REQUIRED_WTS_COLUMNS if col.get("row_required")]
    skipped_incomplete = 0
    eligible_rows: set[int] = set()
    for row, obj in products_by_row.items():
        missing_cells = []
        for key in required_value_keys:
            header_name = resolved_headers.get(key, "")
            if not header_name or to_text(obj.get(header_name, "")) == "":
                missing_cells.append(key)
        if missing_cells:
            skipped_incomplete += 1
            continue
        eligible_rows.add(row)

    if skipped_incomplete:
        log(f"Skipped {skipped_incomplete} row(s) with empty ProdCode, Barcode, or Product Description.")
    if not eligible_rows:
        raise RuntimeError("No product rows left after required value filters.")

    products = []
    for row, obj in products_by_row.items():
        if row not in eligible_rows:
            continue

        product_code = to_text(obj.get(product_code_header, ""))
        barcode = to_text(obj.get(barcode_header, ""))
        product_id = f"{barcode}_{product_code}"

        out = {
            "product_code": product_code,
            "barcode": barcode,
            "product_id": product_id,
            "case_size": max(1, to_int_or_default(obj.get(case_size_header, 1), 1)),
            "minimum_quantity": max(1, to_int_or_default(obj.get(case_size_header, 1), 1)),
            "name": to_text(obj.get(name_header, "")),
            "price": to_float_or_default(obj.get(price_header, 0), 0.0),
            "available_units": to_int_or_default(obj.get(available_header, 0), 0),
            "source": "excel",
            "hazardous": False,
        }

        for optional_key in ("pack_price", "outer_barcode", "outer_per_plt", "layer_qty"):
            header_name = resolved_headers.get(optional_key, "")
            if header_name:
                out[optional_key] = sanitize_cell_text(obj.get(header_name, ""))

        products.append(out)

    payload = {
        "meta": {
            "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            "sourceFile": os.path.basename(excel_path),
            "sheet": sh.title,
            "count": len(products),
            "headerRow": header_row,
            "productIdRule": "product_id = barcode + '_' + product_code",
            "minimumQuantityRule": "minimum_quantity = Pack column",
        },
        "products": products,
    }

    with open(out_json_path, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)

    manifest = {
        "version": datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S"),
        "updated_at": payload["meta"]["generatedAt"],
        "count": len(products),
    }
    with open(DEFAULT_OUT_MANIFEST, "w", encoding="utf-8") as handle:
        json.dump(manifest, handle, ensure_ascii=False, indent=2)

    wb.close()
    t_total = time.perf_counter() - t0
    log("\nDone")
    log(f"- Products: {len(products)}")
    log(f"- JSON written: {out_json_path}")
    log(f"- Manifest written: {DEFAULT_OUT_MANIFEST}")
    log(f"- Total time: {format_duration(t_total)}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
